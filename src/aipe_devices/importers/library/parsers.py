"""Transient file adapters. They never copy uploads into the repository."""

import csv
import io
import json
from dataclasses import dataclass
from zipfile import BadZipFile

from .mapping import known_header
from .models import DetectedTable


@dataclass
class ParsedTable:
    info: DetectedTable
    rows: list[tuple[int, dict]]


def table_from_rows(name, rows, overrides=None):
    rows = list(rows)
    overrides = overrides or {}
    scored = [
        sum(
            known_header(c) or str(c) in overrides or f"{name}.{c}" in overrides
            for c in row
            if c is not None
        )
        for row in rows[:25]
    ]
    header = max(range(len(scored)), key=scored.__getitem__) if scored else 0
    if not rows:
        return ParsedTable(DetectedTable(name=name, header_row=1, columns=(), row_count=0), [])
    columns = tuple(
        str(c).strip() if c is not None else f"__blank_{i}" for i, c in enumerate(rows[header])
    )
    if len(columns) != len(set(columns)):
        raise ValueError(f"{name}: duplicate column names")
    data = []
    for number, row in enumerate(rows[header + 1 :], header + 2):
        if all(v is None or v == "" for v in row):
            continue
        if len(row) > len(columns) and any(v not in (None, "") for v in row[len(columns) :]):
            raise ValueError(f"{name} row {number}: values beyond header")
        data.append((number, dict(zip(columns, row))))
    populated = tuple(c for c in columns if any(r.get(c) not in (None, "") for _, r in data))
    return ParsedTable(
        DetectedTable(name=name, header_row=header + 1, columns=populated, row_count=len(data)),
        data,
    )


class CsvImporter:
    def parse(self, content, overrides=None):
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [table_from_rows("csv", csv.reader(io.StringIO(text), dialect), overrides)]


class ExcelImporter:
    def parse(self, content, overrides=None):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("Excel support requires pip install aipe-devices[library]") from exc
        # Formula text remains visible so an unevaluated/stale cache is never accepted as data.
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except BadZipFile as exc:
            raise ValueError("Invalid Excel workbook") from exc
        try:
            return [
                table_from_rows(sheet.title, sheet.iter_rows(values_only=True), overrides)
                for sheet in book
            ]
        finally:
            book.close()


class GenericJsonImporter:
    def parse(self, content, overrides=None):
        data = json.loads(content)
        if isinstance(data, list):
            data = {"json": data}
        elif isinstance(data, dict) and all(not isinstance(v, (list, dict)) for v in data.values()):
            data = {"json": [data]}
        if not isinstance(data, dict):
            raise ValueError("Generic JSON must be an object or a list of row objects")
        tables = []
        for name, rows in data.items():
            if isinstance(rows, dict):
                rows = [rows]
            if not isinstance(rows, list) or not all(isinstance(r, dict) for r in rows):
                raise ValueError(f"{name}: expected a table of flat row objects")
            if any(isinstance(v, (dict, list)) for r in rows for v in r.values()):
                raise ValueError(f"{name}: nested values require explicit flattening")
            columns = list(dict.fromkeys(k for row in rows for k in row))
            tables.append(
                table_from_rows(
                    name, [columns] + [[row.get(c) for c in columns] for row in rows], overrides
                )
            )
        return tables


class CanonicalJsonImporter:
    def load(self, content):
        from aipe_devices.domain.device import PowerSemiconductorDevice
        from aipe_devices.services.validation import validate_device

        device = PowerSemiconductorDevice.model_validate_json(content)
        return device, validate_device(device)

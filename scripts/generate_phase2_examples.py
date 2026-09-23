"""Generate tiny explicitly synthetic examples and actual Phase-2 API outputs."""

import json
from itertools import product
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from aipe_devices.domain.device import PowerSemiconductorDevice
from aipe_devices.domain.measurement import MeasurementProtocol
from aipe_devices.explorer import DeviceExplorer
from aipe_devices.importers.library import LibraryImporter
from aipe_devices.importers.library.models import ImportReport, ImportSession
from aipe_devices.measurement.gap_planner import draft_measurement_request
from aipe_devices.measurement.plans import SweepAxis, SweepRange, TestPlanItem
from aipe_devices.measurement.request_generator import MeasurementRequest
from aipe_devices.measurement.submission_validator import MeasurementPackage
from aipe_devices.repository.json_repository import JsonDeviceRepository
from aipe_devices.services.coverage import CoverageAnalyzer
from aipe_devices.services.coverage.gaps import CoverageGapReport, DABLossEvaluationPolicy
from aipe_devices.visualization import VisualizationBuilder, VisualizationSpec

ROOT = Path(__file__).resolve().parents[1]


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    text = (
        value.model_dump_json(indent=2)
        if hasattr(value, "model_dump_json")
        else json.dumps(value, indent=2)
    )
    path.write_text(text + "\n", encoding="utf-8")


def main():
    folder = ROOT / "examples/import"
    folder.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    identity = book.active
    identity.title = "identity"
    identity.append(["manufacturer", "part_number", "device_id", "technology"])
    identity.append(["Synthetic Example", "DEMO-001", "synthetic_demo_001", "SiC_MOSFET"])
    switching = book.create_sheet("switching")
    switching.append(["SYNTHETIC tutorial values only; not manufacturer facts or measurements"])
    switching.append(["Tj(C)", "Vdc(V)", "Id(A)", "Rgon(Ohm)", "Rgoff(Ohm)", "Eon(mJ)", "Eoff(mJ)"])
    for tj, voltage, current, resistance in product((25, 125), (600, 800), (10, 20, 40), (2, 5)):
        # Demonstration formula has no physical validation claim.
        eon = current / 40 + voltage / 8000 + resistance / 100 + tj / 1000
        switching.append([tj, voltage, current, resistance, 5, eon, eon / 2])
    conduction = book.create_sheet("conduction")
    conduction.append(["Tj(C)", "Id(A)", "Vds(V)"])
    for tj, current in product((25, 125), (10, 20, 40)):
        conduction.append([tj, current, current * (0.01 + tj / 10000)])
    book.save(folder / "device_switching.xlsx")
    book.close()
    (folder / "device_partial.csv").write_text(
        "manufacturer,part_number,Id(A),Vds(V)\nSynthetic Example,PARTIAL-001,10,0.2\nSynthetic Example,PARTIAL-001,20,0.4\n",
        encoding="utf-8",
    )
    (folder / "device_ambiguous.csv").write_text(
        "manufacturer,part_number,Id(A),Loss1,Loss2\nSynthetic Example,AMBIGUOUS-001,10,0.2,0.1\nSynthetic Example,AMBIGUOUS-001,20,0.4,0.2\n",
        encoding="utf-8",
    )
    write(
        folder / "mapping_override.json",
        {"Loss1": {"semantic": "Eon", "unit": "mJ"}, "Loss2": {"semantic": "Eoff", "unit": "mJ"}},
    )
    write(
        folder / "device_generic.json",
        {
            "identity": {"manufacturer": "Synthetic Example", "part_number": "JSON-001"},
            "conduction": [
                {"Tj(C)": 25, "Id(A)": 10, "Vds(V)": 0.2},
                {"Tj(C)": 25, "Id(A)": 20, "Vds(V)": 0.4},
            ],
        },
    )
    service = LibraryImporter()
    session = service.inspect(folder / "device_switching.xlsx", origin="synthetic")
    if session.report.errors or session.unresolved_fields:
        raise RuntimeError(session.report.model_dump_json())
    device = session.canonical_candidate
    write(folder / "device_canonical.json", device)
    output = ROOT / "examples/phase2"
    write(output / "import_report.json", session.report)
    write(
        output / "ambiguous_report.json",
        service.inspect(folder / "device_ambiguous.csv", origin="synthetic").report,
    )
    dab = CoverageAnalyzer().gaps(device, DABLossEvaluationPolicy())
    write(output / "coverage_gap_report.json", dab)
    write(output / "draft_measurement_request.json", draft_measurement_request(dab))
    write(
        output / "compact_measurement_request.json",
        MeasurementRequest(
            id="synthetic-compact",
            device_id=device.device_id,
            test_plan_items=(
                TestPlanItem(
                    id="dpt",
                    protocol=MeasurementProtocol(id="AIPE-DPT", version="1.0", category="dpt"),
                    sweep_axes=(
                        SweepAxis(
                            parameter="junction_temperature",
                            values=(298.15, 398.15, 423.15),
                            unit="K",
                        ),
                        SweepAxis(parameter="dc_bus_voltage", values=(400, 600, 800), unit="V"),
                        SweepAxis(parameter="drain_current", values=(10, 20, 40, 60), unit="A"),
                        SweepAxis(parameter="gate_resistance_on", values=(2, 5, 10), unit="Ohm"),
                        SweepAxis(parameter="gate_resistance_off", values=(2, 5, 10), unit="Ohm"),
                    ),
                    requested_metrics=("Eon", "Eoff", "dv_dt", "di_dt"),
                    repetitions=3,
                ),
                TestPlanItem(
                    id="cap",
                    protocol=MeasurementProtocol(
                        id="AIPE-CAP", version="1.0", category="capacitance"
                    ),
                    sweep_axes=(
                        SweepAxis(
                            parameter="drain_source_voltage",
                            unit="V",
                            range=SweepRange(start=0, stop=800, step=10),
                        ),
                        SweepAxis(
                            parameter="junction_temperature", unit="K", values=(298.15, 398.15)
                        ),
                    ),
                    requested_metrics=("ciss", "coss", "crss"),
                ),
            ),
        ),
    )
    builder = VisualizationBuilder(
        device_id=device.device_id, device_revision=device.revision, sources=device.provenance
    )
    on, off = device.switching.curves
    fixed = {"dc_bus_voltage": 800, "gate_resistance_on": 5}
    surface = builder.surface(on, x="current", y="junction_temperature", fixed=fixed)
    off_surface = builder.surface(off, x="current", y="junction_temperature", fixed=fixed)
    combined = builder.overlay(surface, off_surface)
    combined = builder.select_slice(
        combined, series_id=on.id, parameter="junction_temperature", value=398.15
    )
    combined = builder.select_point(
        combined, series_id=on.id, coordinates={"current": 40, "junction_temperature": 398.15}
    )
    write(output / "multi_surface.json", combined)
    write(
        output / "line.json",
        builder.line(on, x="current", fixed={**fixed, "junction_temperature": 398.15}),
    )
    write(output / "surface.json", surface)
    write(output / "conduction_surface.json", builder.auto(device.conduction.curves[0]))
    package = MeasurementPackage.model_validate_json(
        (ROOT / "examples/measurement_package/manifest.json").read_bytes()
    )
    write(
        output / "waveform.json",
        builder.waveform(package.runs[0].waveform_refs[0], run=package.runs[0]),
    )
    with TemporaryDirectory() as temp:
        repo = JsonDeviceRepository(temp)
        saved = service.save(session, repo)
        assert saved.status == "imported"
        assert all(p.suffix == ".json" for p in repo.root.iterdir())
        view = DeviceExplorer(repo)
        write(output / "explorer_summary.json", view.summary(device.device_id))
    for name, model in [
        ("import-report", ImportReport),
        ("import-session", ImportSession),
        ("coverage-gap-report", CoverageGapReport),
        ("visualization-spec", VisualizationSpec),
        ("device", PowerSemiconductorDevice),
    ]:
        schema = model.model_json_schema()
        schema["$schema"] = "https://json-schema.org/draft/2020-12/schema"
        write(ROOT / f"docs/schemas/{name}.schema.json", schema)
    print(f"Generated synthetic Phase-2 examples: {folder}")


if __name__ == "__main__":
    main()
